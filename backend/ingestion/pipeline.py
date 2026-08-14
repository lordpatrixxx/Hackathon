import csv
import glob
import hashlib
import json
import logging
import os
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from langchain_text_splitters import RecursiveCharacterTextSplitter

from backend.config import settings

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
    ".csv",
    ".xlsx",
    ".xls",
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
    ".json",
}

# Global in-memory cache for company and index mappings
_COMPANY_MAP: Dict[str, Dict[str, str]] = {}
_INDICES_MAP: Dict[str, str] = {}


def _load_company_mappings(data_dir: Path) -> None:
    global _COMPANY_MAP, _INDICES_MAP
    if _COMPANY_MAP:
        return

    # Load active_companies_list.csv
    active_path = data_dir / "active_companies_list.csv"
    if active_path.exists():
        try:
            with active_path.open("r", encoding="utf-8", errors="ignore") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    scrip = str(row.get("ScripCode", "")).strip()
                    if scrip:
                        _COMPANY_MAP[scrip] = {
                            "name": row.get("CompanyName", "").strip(),
                            "industry": row.get("Industry", "").strip(),
                            "security_id": row.get("SecurityID", "").strip(),
                        }
        except Exception as e:
            logger.warning(f"Error loading active_companies_list: {e}")

    # Load stk.json
    stk_path = data_dir / "stk.json"
    if stk_path.exists():
        try:
            with stk_path.open("r", encoding="utf-8", errors="ignore") as f:
                data = json.load(f)
                for scrip, name in data.items():
                    scrip = str(scrip).strip()
                    if scrip not in _COMPANY_MAP:
                        _COMPANY_MAP[scrip] = {
                            "name": str(name).strip(),
                            "industry": "",
                            "security_id": "",
                        }
        except Exception as e:
            logger.warning(f"Error loading stk.json: {e}")

    # Load indices.json
    indices_path = data_dir / "indices.json"
    if indices_path.exists():
        try:
            with indices_path.open("r", encoding="utf-8", errors="ignore") as f:
                _INDICES_MAP = json.load(f)
        except Exception as e:
            logger.warning(f"Error loading indices.json: {e}")


def get_company_info(scrip_code: str) -> Dict[str, str]:
    scrip_code = str(scrip_code).strip()
    return _COMPANY_MAP.get(
        scrip_code,
        {"name": f"Company (Scrip {scrip_code})", "industry": "", "security_id": ""},
    )


def discover_documents(data_dir: str) -> List[Path]:
    root = Path(data_dir)
    if not root.exists():
        return []
    
    files = []
    # Collect top-level data files first
    for path in sorted(root.glob("*.*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            files.append(path)
            
    # Then collect subdirectories
    for sub in ["stocks", "etfs"]:
        sub_dir = root / sub
        if sub_dir.exists() and sub_dir.is_dir():
            files.append(sub_dir)

    return files


def _hash_chunk(source: str, page: int, text: str, index: int) -> str:
    raw = f"{source}:{page}:{index}:{text}".encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


# -------------------------------------------------------------
# Specialized Semantic Extractors
# -------------------------------------------------------------

def _extract_bse_financial_statements(path: Path) -> List[Dict[str, Any]]:
    docs = []
    name = path.name
    is_bs = "BS_Fin_Stmt" in name
    is_cf = "CF_Fin_Stmt" in name
    is_quarter = "Quarter_Fin_Stmt" in name
    is_yoy = "YOY_Fin_Stmt" in name
    doc_type = "balance_sheet" if is_bs else "cash_flow" if is_cf else "quarterly_results" if is_quarter else "annual_results" if is_yoy else "financial_statement"

    try:
        with path.open("r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            for idx, row in enumerate(reader):
                scrip = str(row.get("ScripCode", "") or row.get("scripCode", "")).strip()
                if not scrip:
                    for k in ["500325", "Scrip", "scrip"]:
                        if k in row:
                            scrip = str(row[k]).strip()
                comp = get_company_info(scrip)
                comp_name = comp["name"] or f"Scrip {scrip}"
                industry = comp["industry"]

                period = row.get("quarter(10yrs)", "") or row.get("quarter", "") or row.get("Mar 2023", "") or ""
                parts = [f"Company: {comp_name} (ScripCode: {scrip})"]
                if industry:
                    parts.append(f"Industry: {industry}")
                parts.append(f"Statement Type: {name} ({doc_type})")
                if period:
                    parts.append(f"Reporting Period: {period}")

                metrics = []
                for k, v in row.items():
                    if not k or k.startswith("Unnamed") or k in {"", "ScripCode", "scripCode", "quarter(10yrs)", "quarter"}:
                        continue
                    v_str = str(v).strip()
                    if v_str and v_str != "-":
                        clean_k = k.replace("\xa0+", "").replace("+", "").strip()
                        metrics.append(f"{clean_k}: {v_str}")

                if metrics:
                    parts.append("Financial Metrics: " + ", ".join(metrics))
                    full_text = " | ".join(parts)
                    docs.append({
                        "text": full_text,
                        "source": str(path),
                        "page": idx + 1,
                        "metadata": {
                            "file_name": name,
                            "document_type": doc_type,
                            "scrip_code": scrip,
                            "company_name": comp_name,
                            "period": period,
                        },
                        "content_type": "table",
                        "extraction_method": "finance_statement_extractor",
                    })
    except Exception as exc:
        logger.warning(f"Error parsing {path.name}: {exc}")
    return docs


def _extract_statement_analysis(path: Path) -> List[Dict[str, Any]]:
    docs = []
    try:
        with path.open("r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            for idx, row in enumerate(reader):
                scrip = str(row.get("ScripCode", "")).strip()
                sentence = str(row.get("sentence", "")).strip()
                header = str(row.get("header", "")).strip()
                dir_status = str(row.get("dir_status", "")).strip()
                comp = get_company_info(scrip)
                comp_name = comp["name"] or f"Scrip {scrip}"

                if sentence:
                    text = f"Company: {comp_name} (ScripCode: {scrip}) | Category: {header} | Analysis: {sentence} | Growth Trend: {dir_status}"
                    docs.append({
                        "text": text,
                        "source": str(path),
                        "page": idx + 1,
                        "metadata": {
                            "file_name": path.name,
                            "document_type": "statement_analysis",
                            "scrip_code": scrip,
                            "company_name": comp_name,
                            "category": header,
                        },
                        "content_type": "text",
                        "extraction_method": "statement_analysis_extractor",
                    })
    except Exception as exc:
        logger.warning(f"Error extracting {path.name}: {exc}")
    return docs


def _extract_corporate_news(path: Path) -> List[Dict[str, Any]]:
    docs = []
    try:
        with path.open("r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            for idx, row in enumerate(reader):
                desc = str(row.get("description", "")).strip()
                date_val = str(row.get("entts", "")).strip()
                newsid = str(row.get("newsid", "")).strip()
                scrip = str(row.get("ScripCode", "") or row.get("scripCode", "")).strip()
                comp = get_company_info(scrip)
                comp_name = comp["name"]

                if desc:
                    text = f"Corporate News Announcement | Company: {comp_name} (Scrip: {scrip}) | Date: {date_val} | News ID: {newsid}\nAnnouncement Details: {desc}"
                    docs.append({
                        "text": text,
                        "source": str(path),
                        "page": idx + 1,
                        "metadata": {
                            "file_name": path.name,
                            "document_type": "corporate_news",
                            "company_name": comp_name,
                            "scrip_code": scrip,
                            "date": date_val,
                        },
                        "content_type": "text",
                        "extraction_method": "news_extractor",
                    })
    except Exception as exc:
        logger.warning(f"Error parsing news {path.name}: {exc}")
    return docs


def _extract_quote_and_peers(path: Path) -> List[Dict[str, Any]]:
    docs = []
    name = path.name
    try:
        with path.open("r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            for idx, row in enumerate(reader):
                scrip = str(row.get("scripCode", "") or row.get("ScripCode", "") or row.get("peers_ScripCode", "")).strip()
                comp = get_company_info(scrip)
                comp_name = row.get("companyName", "") or row.get("peers_Name", "") or comp["name"]
                
                parts = [f"Company: {comp_name} (ScripCode: {scrip})", f"File: {name}"]
                for k, v in row.items():
                    if k and not k.startswith("Unnamed") and k not in {"scripCode", "ScripCode", "peers_ScripCode", "companyName", "peers_Name"}:
                        v_str = str(v).strip()
                        if v_str:
                            parts.append(f"{k}: {v_str}")
                
                text = " | ".join(parts)
                docs.append({
                    "text": text,
                    "source": str(path),
                    "page": idx + 1,
                    "metadata": {
                        "file_name": name,
                        "document_type": "market_quotes_and_peers",
                        "company_name": comp_name,
                        "scrip_code": scrip,
                    },
                    "content_type": "table",
                    "extraction_method": "quote_extractor",
                })
    except Exception as exc:
        logger.warning(f"Error parsing {path.name}: {exc}")
    return docs


def _extract_symbols_valid_meta(path: Path) -> List[Dict[str, Any]]:
    docs = []
    try:
        with path.open("r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            for idx, row in enumerate(reader):
                symbol = str(row.get("Symbol", "")).strip()
                sec_name = str(row.get("Security Name", "")).strip()
                exchange = str(row.get("Listing Exchange", "")).strip()
                is_etf = str(row.get("ETF", "")).strip()
                market_cat = str(row.get("Market Category", "")).strip()

                if symbol:
                    text = f"US Security: {symbol} | Name: {sec_name} | Exchange: {exchange} | Is ETF: {is_etf} | Market Category: {market_cat}"
                    docs.append({
                        "text": text,
                        "source": str(path),
                        "page": idx + 1,
                        "metadata": {
                            "file_name": path.name,
                            "document_type": "security_master",
                            "symbol": symbol,
                            "security_name": sec_name,
                            "exchange": exchange,
                            "is_etf": is_etf,
                        },
                        "content_type": "table",
                        "extraction_method": "symbol_meta_extractor",
                    })
    except Exception as exc:
        logger.warning(f"Error parsing symbols meta: {exc}")
    return docs


def _extract_users_and_cards(path: Path) -> List[Dict[str, Any]]:
    docs = []
    name = path.name
    try:
        with path.open("r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            if "Person" in reader.fieldnames:
                for idx, row in enumerate(reader):
                    person = row.get("Person", "")
                    age = row.get("Current Age", "")
                    ret_age = row.get("Retirement Age", "")
                    gender = row.get("Gender", "")
                    city = row.get("City", "")
                    state = row.get("State", "")
                    income = row.get("Yearly Income - Person", "") or row.get("Income", "")
                    debt = row.get("Total Debt", "")
                    score = row.get("FICO Score", "") or row.get("Credit Score", "")
                    cards = row.get("Num Credit Cards", "")
                    text = f"User Profile: {person} | Age: {age} (Retirement: {ret_age}) | Gender: {gender} | Location: {city}, {state} | Income: {income} | Debt: {debt} | Credit Score: {score} | Cards Count: {cards}"
                    docs.append({
                        "text": text,
                        "source": str(path),
                        "page": idx + 1,
                        "metadata": {
                            "file_name": name,
                            "document_type": "user_profile",
                            "user_name": person,
                            "city": city,
                            "state": state,
                        },
                        "content_type": "table",
                        "extraction_method": "user_extractor",
                    })
            else:
                # Group cards by user
                user_cards = defaultdict(list)
                for row in reader:
                    uid = row.get("User", "")
                    brand = row.get("Card Brand", "")
                    c_type = row.get("Card Type", "")
                    c_num = row.get("Card Number", "")
                    limit = row.get("Credit Limit", "")
                    chip = row.get("Has Chip", "")
                    exp = row.get("Expires", "")
                    user_cards[uid].append(f"{brand} {c_type} (Limit: {limit}, Card: {c_num}, Exp: {exp}, Chip: {chip})")

                for idx, (uid, card_list) in enumerate(user_cards.items(), start=1):
                    text = f"User {uid} Credit Cards Record: Total Cards: {len(card_list)} | " + "; ".join(card_list)
                    docs.append({
                        "text": text,
                        "source": str(path),
                        "page": idx,
                        "metadata": {
                            "file_name": name,
                            "document_type": "card_profile",
                            "user_id": str(uid),
                        },
                        "content_type": "table",
                        "extraction_method": "card_extractor",
                    })
    except Exception as exc:
        logger.warning(f"Error parsing users/cards: {exc}")
    return docs


def _extract_fraud_samples(path: Path, max_rows: int = 500) -> List[Dict[str, Any]]:
    docs = []
    try:
        with path.open("r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            fraud_count = 0
            normal_count = 0
            for idx, row in enumerate(reader):
                is_fraud = str(row.get("is_fraud", "") or row.get("Is Fraud?", "") or row.get("fraud", "")).strip()
                if is_fraud in {"1", "Yes", "true", "TRUE"}:
                    fraud_count += 1
                else:
                    normal_count += 1
                    if normal_count > max_rows:
                        continue

                merchant = row.get("merchant", "") or row.get("Merchant Name", "")
                cat = row.get("category", "")
                amt = row.get("amt", "") or row.get("Amount", "")
                dt = row.get("trans_date_trans_time", "") or f"{row.get('Year','')}-{row.get('Month','')}-{row.get('Day','')}"
                city = row.get("city", "") or row.get("Merchant City", "")
                state = row.get("state", "") or row.get("Merchant State", "")
                
                text = f"Transaction Record ({path.name}) | Date: {dt} | Merchant: {merchant} | Category: {cat} | Amount: {amt} | Location: {city}, {state} | Is Fraud: {is_fraud}"
                docs.append({
                    "text": text,
                    "source": str(path),
                    "page": idx + 1,
                    "metadata": {
                        "file_name": path.name,
                        "document_type": "fraud_transaction",
                        "merchant": str(merchant),
                        "category": str(cat),
                        "is_fraud": is_fraud,
                    },
                    "content_type": "table",
                    "extraction_method": "fraud_extractor",
                })
                if len(docs) >= max_rows * 2:
                    break
    except Exception as exc:
        logger.warning(f"Error parsing fraud samples from {path.name}: {exc}")
    return docs


def _fast_summarize_ticker_file(path: Path) -> Optional[Dict[str, Any]]:
    """Ultra-fast byte-seeking summary of OHLCV historical CSV."""
    ticker = path.stem
    doc_type = "etf_history" if "etfs" in str(path) else "stock_history"
    try:
        with path.open("rb") as f:
            first_line = f.readline().decode("utf-8", "ignore").strip()
            second_line = f.readline().decode("utf-8", "ignore").strip()
            if not first_line or not second_line:
                return None

            # Seek to near end to get last line
            f.seek(0, 2)
            sz = f.tell()
            f.seek(max(0, sz - 2048), 0)
            tail = f.read().decode("utf-8", "ignore").splitlines()
            last_line = tail[-1].strip() if tail else ""
            if not last_line and len(tail) > 1:
                last_line = tail[-2].strip()

        start_parts = second_line.split(",")
        last_parts = last_line.split(",") if last_line else []

        start_date = start_parts[0] if len(start_parts) > 0 else "unknown"
        start_close = start_parts[4] if len(start_parts) > 4 else (start_parts[1] if len(start_parts) > 1 else "unknown")

        latest_date = last_parts[0] if len(last_parts) > 0 else "unknown"
        latest_close = last_parts[4] if len(last_parts) > 4 else (last_parts[1] if len(last_parts) > 1 else "unknown")
        latest_vol = last_parts[6] if len(last_parts) > 6 else (last_parts[-1] if len(last_parts) > 2 else "n/a")

        text = (
            f"Security Historical Profile: {ticker} ({doc_type.upper()}) | "
            f"Historical Trading Range: {start_date} to {latest_date} | "
            f"Latest Close Price: ${latest_close} on {latest_date} (Volume: {latest_vol}) | "
            f"Inception / Initial Close Price: ${start_close} on {start_date} | Source: {path.name}"
        )

        return {
            "text": text,
            "source": str(path),
            "page": 1,
            "metadata": {
                "file_name": path.name,
                "document_type": doc_type,
                "symbol": ticker,
                "start_date": start_date,
                "latest_date": latest_date,
                "latest_close": str(latest_close),
            },
            "content_type": "table",
            "extraction_method": "ticker_history_extractor",
        }
    except Exception as exc:
        return None


def _extract_directory_tickers(dir_path: Path) -> List[Dict[str, Any]]:
    files = list(dir_path.glob("*.csv"))
    docs = []
    with ThreadPoolExecutor(max_workers=16) as ex:
        results = ex.map(_fast_summarize_ticker_file, files)
        for r in results:
            if r:
                docs.append(r)
    return docs


# -------------------------------------------------------------
# General Document Extractors
# -------------------------------------------------------------

def _extract_pdf_text(path: Path) -> List[Dict[str, Any]]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        docs: List[Dict[str, Any]] = []
        for page_index, page in enumerate(reader.pages, start=1):
            content = page.extract_text() or ""
            if content.strip():
                docs.append({
                    "text": content.strip(),
                    "source": str(path),
                    "page": page_index,
                    "metadata": {"file_name": path.name, "document_type": "pdf_document"},
                    "content_type": "text",
                    "extraction_method": "pdf_text",
                })
        return docs
    except Exception:
        return []


def _extract_text_file(path: Path) -> List[Dict[str, Any]]:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
        if text.strip():
            return [{
                "text": text.strip(),
                "source": str(path),
                "page": 1,
                "metadata": {"file_name": path.name, "document_type": "text_document"},
                "content_type": "text",
                "extraction_method": "text",
            }]
    except Exception:
        pass
    return []


def _extract_docx_file(path: Path) -> List[Dict[str, Any]]:
    try:
        from docx import Document
        doc = Document(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            return [{
                "text": "\n".join(paragraphs),
                "source": str(path),
                "page": 1,
                "metadata": {"file_name": path.name, "document_type": "document"},
                "content_type": "text",
                "extraction_method": "docx",
            }]
    except Exception:
        pass
    return []


def _extract_excel_file(path: Path) -> List[Dict[str, Any]]:
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                rows.append(" | ".join("" if value is None else str(value) for value in row))
        if rows:
            return [{
                "text": "\n".join(rows),
                "source": str(path),
                "page": 1,
                "metadata": {"file_name": path.name, "document_type": "table"},
                "content_type": "table",
                "extraction_method": "excel",
            }]
    except Exception:
        pass
    return []


def _extract_image_text(path: Path) -> List[Dict[str, Any]]:
    try:
        import pytesseract
        from PIL import Image
        image = Image.open(path)
        text = pytesseract.image_to_string(image)
        if text.strip():
            return [{
                "text": text.strip(),
                "source": str(path),
                "page": 1,
                "metadata": {"file_name": path.name, "document_type": "image"},
                "content_type": "ocr",
                "extraction_method": "ocr",
            }]
    except Exception:
        pass
    return []


# -------------------------------------------------------------
# Main Extraction Router
# -------------------------------------------------------------

def extract_document_text(path: Path) -> List[Dict[str, Any]]:
    data_dir = path.parent if path.is_file() else path
    _load_company_mappings(data_dir if data_dir.name == "data" else data_dir.parent)

    if path.is_dir():
        if path.name in {"stocks", "etfs"}:
            return _extract_directory_tickers(path)
        return []

    name = path.name.lower()
    suffix = path.suffix.lower()

    if "fin_stmt" in name:
        return _extract_bse_financial_statements(path)
    if "statement_analysis" in name:
        return _extract_statement_analysis(path)
    if "corporate_news" in name:
        return _extract_corporate_news(path)
    if name in {"quote_data.csv", "quote_data_buy.csv", "quote_data_sell.csv", "peers_comparisons_data.csv"}:
        return _extract_quote_and_peers(path)
    if name == "symbols_valid_meta.csv":
        return _extract_symbols_valid_meta(path)
    if name in {"sd254_users.csv", "sd254_cards.csv"}:
        return _extract_users_and_cards(path)
    if "fraud" in name or name in {"cc_fraud.csv", "user0_credit_card_transactions.csv"}:
        return _extract_fraud_samples(path)
    if name == "stk.json":
        # Master company lookup
        try:
            with path.open("r", encoding="utf-8", errors="ignore") as f:
                data = json.load(f)
            docs = []
            items = list(data.items())
            batch_sz = 30
            for i in range(0, len(items), batch_sz):
                batch = items[i : i + batch_sz]
                listing = ", ".join(f"{name} (Scrip: {code})" for code, name in batch)
                docs.append({
                    "text": f"BSE Listed Companies Master Directory (Batch {i//batch_sz + 1}): {listing}",
                    "source": str(path),
                    "page": i // batch_sz + 1,
                    "metadata": {"file_name": path.name, "document_type": "bse_company_directory"},
                    "content_type": "text",
                    "extraction_method": "json_directory",
                })
            return docs
        except Exception:
            return []

    if name in {"indices.json", "indexes_df.csv"}:
        docs = []
        if name == "indices.json":
            try:
                with path.open("r", encoding="utf-8", errors="ignore") as f:
                    data = json.load(f)
                    for code, title in data.items():
                        docs.append({
                            "text": f"Index Reference: Code {code} -> Name: {title}",
                            "source": str(path),
                            "page": 1,
                            "metadata": {"file_name": path.name, "document_type": "index_master", "index_code": code, "index_name": title},
                            "content_type": "text",
                            "extraction_method": "json",
                        })
            except Exception:
                pass
        return docs

    # Large multi-GB dataset statistical schema summaries
    if name in {"credit_card_transactions-ibm_v2.csv", "credit_card_transactions.csv", "stocks_df.csv", "transactions_df.csv", "nse_indexes.csv"}:
        try:
            with path.open("r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                header = next(reader, [])
                sample_rows = [next(reader, []) for _ in range(5)]
                
            sample_str = "\n".join(" | ".join(r) for r in sample_rows if r)
            text = (
                f"Large Dataset Master Summary: {path.name} | File Size: {path.stat().st_size / (1024*1024):.2f} MB | "
                f"Columns ({len(header)}): {', '.join(header)}\n"
                f"Sample Records:\n{sample_str}"
            )
            return [{
                "text": text,
                "source": str(path),
                "page": 1,
                "metadata": {"file_name": path.name, "document_type": "large_dataset_master", "columns": ", ".join(header[:10])},
                "content_type": "table",
                "extraction_method": "large_file_summary",
            }]
        except Exception:
            return []

    # Standard formats
    if suffix == ".pdf":
        return _extract_pdf_text(path)
    if suffix == ".txt":
        return _extract_text_file(path)
    if suffix == ".docx":
        return _extract_docx_file(path)
    if suffix in {".xlsx", ".xls"}:
        return _extract_excel_file(path)
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
        return _extract_image_text(path)
    if suffix == ".csv":
        try:
            with path.open("r", encoding="utf-8", errors="ignore") as f:
                reader = csv.DictReader(f)
                docs = []
                for idx, row in enumerate(reader):
                    row_str = " | ".join(f"{k}: {v}" for k, v in row.items() if k and v)
                    if row_str:
                        docs.append({
                            "text": f"File: {path.name} | {row_str}",
                            "source": str(path),
                            "page": idx + 1,
                            "metadata": {"file_name": path.name, "document_type": "table"},
                            "content_type": "table",
                            "extraction_method": "generic_csv",
                        })
                    if len(docs) >= 1000:
                        break
                return docs
        except Exception:
            return []
    return []


def chunk_documents(documents: Iterable[Dict[str, Any]], chunk_size: int = None, chunk_overlap: int = None) -> List[Dict[str, Any]]:
    chunk_size = chunk_size or settings.CHUNK_SIZE
    chunk_overlap = chunk_overlap or settings.CHUNK_OVERLAP
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " | ", " ", ""],
    )
    chunks: List[Dict[str, Any]] = []
    for entry in documents:
        text = (entry.get("text") or "").strip()
        if not text:
            continue
        
        if len(text) <= chunk_size:
            split = [text]
        else:
            split = splitter.split_text(text)

        for index, part in enumerate(split):
            if not part.strip():
                continue
            metadata = dict(entry.get("metadata") or {})
            page = entry.get("page")
            source = entry.get("source")
            chunk_id = _hash_chunk(source, page or 1, part, index)
            metadata.update({
                "source": source,
                "file_name": Path(source).name if source else "unknown",
                "page": page or 1,
                "chunk_id": chunk_id,
                "content_type": entry.get("content_type", "text"),
                "extraction_method": entry.get("extraction_method", "text"),
            })
            chunks.append({
                "text": part,
                "source": source,
                "page": page or 1,
                "metadata": metadata,
                "content_type": entry.get("content_type", "text"),
                "extraction_method": entry.get("extraction_method", "text"),
                "chunk_id": chunk_id,
            })
    return chunks
